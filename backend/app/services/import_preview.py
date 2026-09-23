"""Parse first, recheck conflicts under a database lock, then commit atomically."""

import json
from datetime import time

from openpyxl import load_workbook
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import select, text

from app.models import Classroom, Discipline, Group, Schedule, Teacher, WeekType
from app.services import schedule_import, timetable_import
from app.services.file_validation import validate_workbook


class Lesson(BaseModel):
    group: str = Field(min_length=1, max_length=50)
    discipline: str = Field(min_length=1, max_length=300)
    teacher: str | None = Field(default=None, max_length=200)
    classroom: str | None = Field(default=None, max_length=50)
    weekday: int = Field(ge=1, le=7)
    starts_at: time
    ends_at: time
    week_type: WeekType = WeekType.every
    lesson_type: str | None = None

    @model_validator(mode="after")
    def valid_times(self):
        if self.ends_at <= self.starts_at:
            raise ValueError("Окончание должно быть позже начала")
        return self


def parse(stream):
    validate_workbook(stream)
    if timetable_import.looks_like_timetable(stream):
        stream.seek(0)
        lessons, errors = timetable_import.parse_workbook(stream)
        from dataclasses import asdict
        return [Lesson.model_validate(asdict(l)) for l in lessons], errors
    stream.seek(0)
    wb = load_workbook(stream, read_only=True, data_only=False)
    try:
        rows = wb.active.iter_rows(values_only=True)
        columns = schedule_import._map_headers(next(rows, []))
        lessons, errors = [], []
        for number, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            try:
                values = {key: row[index] for key,index in columns.items()}
                for key in ("group","teacher","discipline","classroom"):
                    values[key] = str(values[key]).strip() if values.get(key) is not None else None
                values["weekday"] = schedule_import._parse_weekday(values["weekday"])
                values["starts_at"] = schedule_import._parse_time(values["starts_at"])
                values["ends_at"] = schedule_import._parse_time(values["ends_at"])
                week = str(values.get("week_type") or "").strip().lower()
                if week not in schedule_import.WEEK_TYPES:
                    raise ValueError("Неизвестный тип недели")
                values["week_type"] = schedule_import.WEEK_TYPES[week]
                lessons.append(Lesson.model_validate(values))
            except (ValueError, TypeError, IndexError) as exc:
                errors.append(f"Строка {number}: некорректное значение ({type(exc).__name__})")
        return lessons, errors
    finally:
        wb.close()


def as_lesson(item):
    return Lesson(group=item.group.name, discipline=item.discipline.name,
        teacher=item.teacher.full_name if item.teacher else None,
        classroom=item.classroom.number if item.classroom else None,
        weekday=item.weekday, starts_at=item.starts_at, ends_at=item.ends_at,
        week_type=item.week_type, lesson_type=item.lesson_type)


def conflicts(db, lessons):
    known = [as_lesson(item) for item in db.scalars(select(Schedule).order_by(Schedule.id))]
    errors, accepted, skipped = [], [], 0
    for index, item in enumerate(lessons, 1):
        if item in known:
            skipped += 1
            continue
        for other in known:
            weeks_overlap = (item.week_type == other.week_type or WeekType.every in {item.week_type, other.week_type})
            times_overlap = item.starts_at < other.ends_at and other.starts_at < item.ends_at
            if item.weekday != other.weekday or not weeks_overlap or not times_overlap:
                continue
            reasons = []
            if item.group == other.group:
                reasons.append("группа")
            if item.teacher and item.teacher == other.teacher:
                reasons.append("преподаватель")
            if item.classroom and item.classroom == other.classroom:
                reasons.append("аудитория")
            if reasons:
                errors.append(f"Запись {index}: пересечение по полям {', '.join(reasons)}")
        known.append(item)
        accepted.append(item)
    return accepted, skipped, errors


def lock_schedule(db):
    if db.bind.dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(78230412)"))


def apply(db, lessons):
    created = 0
    for lesson in lessons:
        get = schedule_import._get_or_create
        group = get(db, Group, {"name": lesson.group})
        discipline = get(db, Discipline, {"name": lesson.discipline})
        teacher = get(db, Teacher, {"full_name":lesson.teacher}) if lesson.teacher else None
        classroom = get(db, Classroom, {"number":lesson.classroom}) if lesson.classroom else None
        db.add(Schedule(group_id=group.id, discipline_id=discipline.id,
            teacher_id=teacher.id if teacher else None, classroom_id=classroom.id if classroom else None,
            weekday=lesson.weekday, starts_at=lesson.starts_at, ends_at=lesson.ends_at,
            week_type=lesson.week_type, lesson_type=lesson.lesson_type))
        created += 1
    db.flush()
    return created
